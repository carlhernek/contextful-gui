"""Tests for file preview (text, csv, xlsx)."""
from __future__ import annotations

from pathlib import Path

import pytest

from contextful_sidecar.runtime.preview import preview_file
from contextful_sidecar.runtime.runs import _list_meta_docs


@pytest.fixture
def workspace(tmp_path: Path) -> Path:
    ws = tmp_path / "ws"
    (ws / "meta").mkdir(parents=True)
    (ws / "meta" / "notes.md").write_text("# Notes\nhello", encoding="utf-8")
    (ws / "meta" / "data.csv").write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    (ws / "meta" / "sub").mkdir()
    (ws / "meta" / "sub" / "nested.md").write_text("nested", encoding="utf-8")
    return ws


def test_preview_text_md(workspace: Path) -> None:
    result = preview_file(workspace, "notes.md", base="meta")
    assert result["ok"] is True
    assert result["kind"] == "text"
    assert "hello" in result["content"]


def test_preview_csv_table(workspace: Path) -> None:
    result = preview_file(workspace, "data.csv", base="meta")
    assert result["ok"] is True
    assert result["kind"] == "table"
    assert result["table"]["headers"] == ["a", "b"]
    assert len(result["table"]["rows"]) == 2


def test_preview_xlsx_table(workspace: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["col1", "col2"])
    ws.append(["x", "y"])
    path = workspace / "meta" / "sheet.xlsx"
    wb.save(path)

    result = preview_file(workspace, "sheet.xlsx", base="meta")
    assert result["ok"] is True
    assert result["kind"] == "table"
    assert result["table"]["headers"] == ["col1", "col2"]


def test_preview_unsupported_binary(workspace: Path) -> None:
    (workspace / "meta" / "image.bin").write_bytes(b"\x00\x01\x02")
    result = preview_file(workspace, "image.bin", base="meta")
    assert result["ok"] is False
    assert result["kind"] == "unsupported"


def test_preview_docx_text(workspace: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello requirements")
    target = workspace / "meta" / "spec.docx"
    doc.save(target)

    result = preview_file(workspace, "spec.docx", base="meta")
    assert result["ok"] is True
    assert "Hello requirements" in result.get("content", "")


def test_preview_png_image(workspace: Path) -> None:
    # minimal 1x1 PNG
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
        b"\x0d\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    (workspace / "meta" / "icon.png").write_bytes(png)
    result = preview_file(workspace, "icon.png", base="meta")
    assert result["ok"] is True
    assert result["kind"] == "image"
    assert result.get("imageUrl", "").startswith("data:image/")


def test_preview_meta_base(workspace: Path) -> None:
    result = preview_file(workspace, "notes.md", base="meta")
    assert result["ok"] is True
    assert "meta/" in result["path"]


def test_list_meta_docs_recursive(workspace: Path) -> None:
    docs = _list_meta_docs(workspace)
    names = {p.name for p in docs}
    assert "notes.md" in names
    assert "nested.md" in names
